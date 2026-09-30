from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from PIL import Image as PILImage, ImageDraw, ImageFont

out=Path(__file__).parent
font='C:/Windows/Fonts/malgun.ttf'
bold='C:/Windows/Fonts/malgunbd.ttf'
pdfmetrics.registerFont(TTFont('Malgun',font)); pdfmetrics.registerFont(TTFont('MalgunBold',bold))
pdfmetrics.registerFontFamily('Malgun',normal='Malgun',bold='MalgunBold')
doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297)
sec.top_margin=Mm(25); sec.bottom_margin=Mm(22); sec.left_margin=sec.right_margin=Mm(24)
sec.header_distance=sec.footer_distance=Mm(10)
for name,size in [('Normal',11),('Title',18),('Heading 1',18),('Heading 2',14)]:
 s=doc.styles[name]; s.font.name='Malgun Gothic'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string('151A20')
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
 s.paragraph_format.line_spacing=1.4; s.paragraph_format.space_after=Mm(3)
 if name.startswith('Heading'): s.paragraph_format.keep_with_next=True; s.paragraph_format.space_before=Mm(9); s.paragraph_format.space_after=Mm(4)
styles={
 'body':ParagraphStyle('body',fontName='Malgun',fontSize=11,leading=15.4,spaceAfter=3*mm,textColor=HexColor('#151A20'),wordWrap='CJK'),
 'h1':ParagraphStyle('h1',fontName='MalgunBold',fontSize=18,leading=25.2,spaceAfter=6*mm,textColor=HexColor('#1A304B'),keepWithNext=True),
 'h2':ParagraphStyle('h2',fontName='MalgunBold',fontSize=14,leading=19.6,spaceBefore=9*mm,spaceAfter=4*mm,textColor=HexColor('#1A304B'),keepWithNext=True),
 'cell':ParagraphStyle('cell',fontName='Malgun',fontSize=10.5,leading=14.7,wordWrap='CJK')}
story=[]
def para(text,style='body'):
 story.append(Paragraph(text.replace('<b>','<b><font color="#245A81">').replace('</b>','</font></b>'),styles[style]))
 p=doc.add_paragraph(style={'body':'Normal','h1':'Heading 1','h2':'Heading 2'}[style])
 # Short bold labels keep meaning readable in grayscale.
 if '<b>' in text:
  first,rest=text.split('</b>',1); r=p.add_run(first.replace('<b>','')); r.bold=True; r.font.color.rgb=RGBColor.from_string('245A81'); p.add_run(rest)
 else: p.add_run(text)
 for r in p.runs:
  r.font.name='Malgun Gothic'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
 return p
def table(rows,widths):
 t=Table([[Paragraph(c,styles['cell']) for c in row] for row in rows],colWidths=[x*mm for x in widths],repeatRows=1)
 t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#EEEEEE')),('GRID',(0,0),(-1,-1),0.5,HexColor('#BDBDBD')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),3*mm),('RIGHTPADDING',(0,0),(-1,-1),3*mm),('TOPPADDING',(0,0),(-1,-1),2.5*mm),('BOTTOMPADDING',(0,0),(-1,-1),2.5*mm)]))
 story.extend([t,Spacer(1,9*mm)])
 dt=doc.add_table(rows=0,cols=len(widths)); dt.autofit=False
 for row in rows:
  cells=dt.add_row().cells
  for c,txt,w in zip(cells,row,widths):
   c.width=Mm(w); c.text=txt
   props=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
   for edge,val in [('top',142),('bottom',142),('left',170),('right',170)]:
    el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); margins.append(el)
   props.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.keep_together=True
    for r in p.runs: r.font.size=Pt(10.5); r.font.name='Malgun Gothic'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
  trpr=cells[0]._tc.getparent().get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
 hdr=dt.rows[0]._tr.get_or_add_trPr(); hdr.append(OxmlElement('w:tblHeader'))
 for c in dt.rows[0].cells:
  sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEEEEE'); c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs:r.bold=True
def page():story.append(PageBreak());doc.add_page_break()
def screen(name,complex=False):
 im=PILImage.new('RGB',(1100,650 if complex else 230),'white'); d=ImageDraw.Draw(im)
 f=ImageFont.truetype(font,27); fb=ImageFont.truetype(bold,30)
 def box(x,y,w,h,txt,fill='white'):
  d.rounded_rectangle((x,y,x+w,y+h),radius=5,outline='#66727E',width=2,fill=fill);d.text((x+16,y+13),txt,font=f,fill='#151A20')
 d.rectangle((2,2,1097,im.height-3),outline='#BDBDBD',width=2)
 if not complex:
  d.text((28,22),'예약 검색',font=fb,fill='#151A20');box(28,85,780,62,'검색어 입력');box(838,85,230,62,'검색','#EEEEEE')
 else:
  d.text((28,22),'예약 정보 및 승인 처리',font=fb,fill='#151A20')
  d.text((28,86),'예약 정보',font=fb,fill='#151A20');box(28,132,500,62,'요청번호 입력');box(558,132,510,62,'날짜 입력')
  box(28,217,270,62,'승인 대기로 저장','#EEEEEE');d.text((328,230),'상태  승인 대기',font=f,fill='#151A20')
  d.text((28,316),'반려 의견  필수',font=fb,fill='#151A20');box(28,365,1040,70,'반려 사유 입력')
  box(28,459,235,62,'승인','#EEEEEE');box(293,459,235,62,'반려','#EEEEEE')
  d.text((28,551),'저장 실패  입력한 내용을 유지합니다',font=f,fill='#151A20');box(838,541,230,62,'재시도','#EEEEEE')
 im.save(out/name)
 return im.height/1100*162
def picture(name,h):
 story.extend([Image(str(out/name),width=162*mm,height=h*mm),Spacer(1,4*mm)])
 doc.add_picture(str(out/name),width=Mm(162)); doc.paragraphs[-1].paragraph_format.space_after=Mm(4)

para('예약 서비스 기능과 화면 구성','h1')
para('예약 생성과 CSV 내보내기는 확정된 개발 범위다. 운영자는 예약을 승인 대기로 저장하고 승인 또는 반려로 처리한다. 중복 판단 시간범위와 보유기간은 결정이 필요하다.')
para('요구사항','h2')
table([['ID','기능','상태'],['REQ-01','예약 생성','확정'],['REQ-02','CSV 내보내기','확정']],[25,110,27])
para('<b>확인사항: REQ-01</b> — 중복 판단에 사용할 시간범위를 결정한다.')
para('<b>확인사항: REQ-02</b> — 내보내기 대상 데이터의 보유기간을 결정한다.')
para('REQ-01 예약 생성','h2')
para('운영자가 날짜와 요청번호를 입력하고 중복 검사 후 승인 대기로 저장한다.')
para('<b>역할과 진입 조건</b> — 운영자가 예약 생성 화면에서 시작한다.')
para('<b>입력과 검증</b> — 날짜와 요청번호를 입력하고 동일 요청번호를 검사한다.')
para('<b>처리와 상태</b>')
para('1. 동일 요청번호를 검사하고 중복 요청을 차단한다.')
para('2. 중복이 없으면 승인 대기로 저장한다.')
para('3. 승인하면 확정된다. 반려할 때는 의견을 반드시 입력한다.')
para('<b>결과와 오류 복구</b> — 중복 요청은 차단한다. 저장 실패 시 입력을 유지하고 재시도를 제공한다.')
para('<b>완료 기준</b> — 중복 차단과 승인 및 반려 이력을 검수한다.')
page()
para('REQ-02 CSV 내보내기','h2')
para('CSV 내보내기를 제공한다. 데이터 보유기간은 확인 후 적용한다.')
para('<b>확인사항: REQ-02</b> — 내보내기 권한, 포함 항목, 오류 처리와 검수 방법은 별도 결정이 필요하다.')
para('기술스택','h2')
table([['분야','기술','용도'],['프론트엔드','React','예약 입력 및 승인 화면'],['백엔드','FastAPI','예약 처리 및 CSV 내보내기'],['인프라','미정','운영 환경 결정 필요']],[36,43,83])
para('검색 화면','h2')
picture('search.png',screen('search.png'))
para('검색어를 입력하고 검색 버튼을 누른다. 검색 대상과 결과 표시 방식은 확인이 필요하다.')
page()
para('예약 처리 화면','h2')
picture('booking.png',screen('booking.png',True))
para('<b>입력과 동작</b> — 요청번호와 날짜를 입력하여 승인 대기로 저장한다. 승인하면 확정된다. 반려하려면 반려 의견을 입력해야 한다.')
para('<b>상태와 인계</b> — 승인 대기에서 승인 또는 반려로 처리한다. 중복 차단 결과와 승인 및 반려 이력을 확인할 수 있어야 한다.')
para('<b>오류 복구</b> — 저장 실패 시 입력한 값을 유지하고 재시도로 저장을 다시 시도한다.')
para('<b>확인사항: REQ-01</b> — 중복 판단 시간범위와 승인 및 반려 이력의 표시 위치를 결정한다.')
doc.save(out/'reservation-sample.docx')
SimpleDocTemplate(str(out/'reservation-sample.pdf'),pagesize=(210*mm,297*mm),topMargin=25*mm,bottomMargin=22*mm,leftMargin=24*mm,rightMargin=24*mm).build(story)
print('Created DOCX and separately authored PDF')
